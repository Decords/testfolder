#!/usr/bin/env python3
"""Build a cumulative PHH/Pigu product XML and price feeds from Shopify bulk JSONL.

Rules implemented for Decords:
- New batch: up to 500 eligible active Wall Decals, excluding already imported products.
- Multi-variant parent titles do not contain a size or size range.
- Exact rounded centimetre size is included in every modification title.
- Single fixed-size products include the exact centimetre size in the parent title.
- Stock is always 999.
- Collection/production time is always 120 hours for LT/LV/EE/FI.
- Current Shopify Markets prices are used separately by country.
- Existing batch 001 product XML is copied unchanged into the cumulative feed.
"""
from __future__ import annotations

import argparse
import csv
import html
import json
import math
import re
import sys
from collections import defaultdict
from copy import deepcopy
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

import requests
from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

LANGS = ("lt", "ru", "lv", "ee", "fi")
TITLE_TAG = {"lt": "title", "ru": "title-ru", "lv": "title-lv", "ee": "title-ee", "fi": "title-fi"}
DESC_TAG = {
    "lt": "long-description",
    "ru": "long-description-ru",
    "lv": "long-description-lv",
    "ee": "long-description-ee",
    "fi": "long-description-fi",
}
MOD_TAG = {
    "lt": "modification-title",
    "ru": "modification-title-ru",
    "lv": "modification-title-lv",
    "ee": "modification-title-ee",
    "fi": "modification-title-fi",
}
TRANSLATION_ALIAS = {
    "lt": "translationsLt",
    "lv": "translationsLv",
    "ee": "translationsEt",
    "fi": "translationsFi",
    "ru": "translationsRu",
}

IMPORTED_PRODUCT_IDS = {
    "7348164264186",  # Words Are Powerful pilot
    "8755085312338",
    "10170330317138",
    "10170349683026",
    "8755067322706",
    "10189851590994",
    "8568369676626",
    "7348059668730",
    "8755081478482",
    "8609989067090",
    "8568466571602",
}

CATEGORY_ID = "1001"
CATEGORY_NAME = "Baldai ir namų interjeras / Namų interjeras / Interjero lipdukai"
MANUFACTURER_NAME = "AnOL OÜ"
MANUFACTURER_ADDRESS = "Laki 4-9, 10621 Tallinn, Estonia"
MANUFACTURER_EMAIL = "info@decords.com"

PRODUCTION_HOURS = 120
STOCK = 999
TARGET_PRODUCTS = 500

INCH_WORDS = r'(?:inches?|inch|in\.?|\"|tolli(?:a|t)?|tuumaa|дюйм(?:а|ов)?|coll(?:i|u)?)'
CM_WORDS = r'(?:cm|см|centimetr(?:as|ai|ų|u|us|i|e)?|centimetrs|sentimeetrit?|centtimetri(?:ä)?)'
PAIR_RE = re.compile(
    rf'(?<!\d)(\d+(?:[\.,]\d+)?)\s*[x×XхХ]\s*(\d+(?:[\.,]\d+)?)\s*({CM_WORDS}|{INCH_WORDS})?',
    re.IGNORECASE,
)
SINGLE_UNIT_RE = re.compile(rf'(?<!\d)(\d+(?:[\.,]\d+)?)\s*({CM_WORDS}|{INCH_WORDS})', re.IGNORECASE)
HTML_TAG_RE = re.compile(r'<[^>]+>')
SPACE_RE = re.compile(r'\s+')

CUSTOM_WORDS = re.compile(
    r'\b(custom|customized|customised|personalized|personalised|made\s+to\s+order\s+with\s+name|savo\s+tekst|персонал|индивидуальн)\b',
    re.IGNORECASE,
)

INSTALL_BLOCKS = {
    "lt": (
        "<h3>Matmenys ir pasirinkimas</h3><p>{sizes}</p>"
        "<h3>Montavimo instrukcija</h3><ol>"
        "<li>Prieš klijuodami išmatuokite paviršių ir suplanuokite galutinę kompoziciją.</li>"
        "<li>Paviršius turi būti lygus, švarus, sausas ir visiškai sukietėjęs. Pašalinkite dulkes, riebalus ir drėgmę.</li>"
        "<li>Prieš visą montavimą išbandykite mažame nepastebimame plote.</li>"
        "<li>Klijuokite palaipsniui nuo centro į kraštus, nestempdami vinilo ir išstumdami orą.</li>"
        "<li>Tvirtai prispauskite kraštus. Pirmąsias 24 valandas lipduko neplaukite ir netrinkite.</li>"
        "</ol><h3>Priežiūra ir nuėmimas</h3>"
        "<p>Valykite minkšta sausa arba vos drėgna šluoste. Nenaudokite abrazyvų ar tirpiklių. Nuimdami lengvai pašildykite ir lėtai traukite mažu kampu.</p>"
    ),
    "ru": (
        "<h3>Размеры и выбор</h3><p>{sizes}</p>"
        "<h3>Инструкция по установке</h3><ol>"
        "<li>Перед нанесением измерьте поверхность и спланируйте итоговую композицию.</li>"
        "<li>Поверхность должна быть гладкой, чистой, сухой и полностью отвердевшей. Удалите пыль, жир и влагу.</li>"
        "<li>Сначала выполните тест на небольшом незаметном участке.</li>"
        "<li>Наклеивайте постепенно от центра к краям, не растягивая винил и выгоняя воздух.</li>"
        "<li>Плотно прижмите края. Не мойте и не трите наклейку первые 24 часа.</li>"
        "</ol><h3>Уход и снятие</h3>"
        "<p>Протирайте мягкой сухой или слегка влажной тканью. Не используйте абразивы и растворители. Для снятия слегка прогрейте изделие и медленно отделяйте под небольшим углом.</p>"
    ),
    "lv": (
        "<h3>Izmēri un izvēle</h3><p>{sizes}</p>"
        "<h3>Uzstādīšanas instrukcija</h3><ol>"
        "<li>Pirms uzstādīšanas izmēriet virsmu un izplānojiet gala kompozīciju.</li>"
        "<li>Virsmai jābūt gludai, tīrai, sausai un pilnībā sacietējušai. Notīriet putekļus, taukus un mitrumu.</li>"
        "<li>Vispirms pārbaudiet nelielā neuzkrītošā vietā.</li>"
        "<li>Līmējiet pakāpeniski no centra uz malām, nestiepjot vinilu un izspiežot gaisu.</li>"
        "<li>Stingri piespiediet malas. Pirmās 24 stundas uzlīmi nemazgājiet un neberzējiet.</li>"
        "</ol><h3>Kopšana un noņemšana</h3>"
        "<p>Tīriet ar mīkstu sausu vai nedaudz mitru drānu. Nelietojiet abrazīvus līdzekļus vai šķīdinātājus. Noņemšanai viegli sasildiet un lēnām velciet nelielā leņķī.</p>"
    ),
    "ee": (
        "<h3>Mõõdud ja valik</h3><p>{sizes}</p>"
        "<h3>Paigaldusjuhend</h3><ol>"
        "<li>Enne paigaldamist mõõtke pind ja kavandage lõplik kompositsioon.</li>"
        "<li>Pind peab olema sile, puhas, kuiv ja täielikult kõvenenud. Eemaldage tolm, rasv ja niiskus.</li>"
        "<li>Tehke esmalt katse väikesel varjatud kohal.</li>"
        "<li>Kleepige järk-järgult keskelt servade suunas, vinüüli venitamata ja õhku välja surudes.</li>"
        "<li>Suruge servad korralikult kinni. Ärge peske ega hõõruge kleebist esimese 24 tunni jooksul.</li>"
        "</ol><h3>Hooldus ja eemaldamine</h3>"
        "<p>Puhastage pehme kuiva või kergelt niiske lapiga. Ärge kasutage abrasiive ega lahusteid. Eemaldamiseks soojendage kergelt ja tõmmake aeglaselt väikese nurga all.</p>"
    ),
    "fi": (
        "<h3>Mitat ja valinta</h3><p>{sizes}</p>"
        "<h3>Asennusohje</h3><ol>"
        "<li>Mittaa pinta ja suunnittele lopullinen sommitelma ennen asennusta.</li>"
        "<li>Pinnan on oltava sileä, puhdas, kuiva ja täysin kovettunut. Poista pöly, rasva ja kosteus.</li>"
        "<li>Testaa ensin pienessä huomaamattomassa kohdassa.</li>"
        "<li>Kiinnitä vähitellen keskeltä reunoja kohti venyttämättä vinyyliä ja poista ilma.</li>"
        "<li>Paina reunat huolellisesti kiinni. Älä pese tai hankaa tarraa ensimmäisten 24 tunnin aikana.</li>"
        "</ol><h3>Hoito ja irrotus</h3>"
        "<p>Pyyhi pehmeällä kuivalla tai hieman kostealla liinalla. Älä käytä hankaavia välineitä tai liuottimia. Irrotusta varten lämmitä kevyesti ja vedä hitaasti pienessä kulmassa.</p>"
    ),
}

SIZE_INTRO = {
    "lt": "Galimos modifikacijos: {items}.",
    "ru": "Доступные модификации: {items}.",
    "lv": "Pieejamās modifikācijas: {items}.",
    "ee": "Saadaval modifikatsioonid: {items}.",
    "fi": "Saatavilla olevat vaihtoehdot: {items}.",
}

FALLBACK_TITLE_SUFFIX = {
    "lt": "sienų lipdukas Decords",
    "ru": "настенная наклейка Decords",
    "lv": "sienas uzlīme Decords",
    "ee": "seinakleebis Decords",
    "fi": "seinätarra Decords",
}


@dataclass
class SizeInfo:
    kind: str  # exact, range, single
    width_cm: int | None = None
    height_cm: int | None = None
    min_cm: int | None = None
    max_cm: int | None = None
    source: str = ""
    assumed_inches: bool = False

    def label(self, lang: str = "lt") -> str:
        unit = "см" if lang == "ru" else "cm"
        if self.kind == "exact" and self.width_cm is not None and self.height_cm is not None:
            return f"{self.width_cm} × {self.height_cm} {unit}"
        if self.kind == "range" and self.min_cm is not None and self.max_cm is not None:
            return f"{self.min_cm}–{self.max_cm} {unit}"
        if self.kind == "single" and self.max_cm is not None:
            return f"{self.max_cm} {unit}"
        return ""


@dataclass
class VariantRecord:
    raw: dict[str, Any]
    sku: str
    ean: str
    size: SizeInfo
    extra_label: str
    prices: dict[str, tuple[Decimal, Decimal]]
    weight_kg: Decimal
    length_m: Decimal
    height_m: Decimal = Decimal("0.10")
    width_m: Decimal = Decimal("0.10")


@dataclass
class ProductRecord:
    raw: dict[str, Any]
    product_id: str
    title: str
    handle: str
    images: list[str]
    variants: list[VariantRecord]
    localized_titles: dict[str, str]
    localized_descriptions: dict[str, str]
    missing_translations: list[str]
    outdated_translations: list[str]
    size_confidence: str
    color_lt: str
    pattern_lt: str


def decimal(value: Any) -> Decimal | None:
    if value is None:
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None


def q2(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def numeric_product_id(gid: str) -> str:
    return gid.rsplit("/", 1)[-1]


def valid_ean13(ean: str) -> bool:
    if not re.fullmatch(r"\d{13}", ean or ""):
        return False
    digits = [int(c) for c in ean]
    check = (10 - ((sum(digits[:-1:2]) + 3 * sum(digits[1:-1:2])) % 10)) % 10
    return check == digits[-1]


def strip_html(text: str) -> str:
    return SPACE_RE.sub(" ", html.unescape(HTML_TAG_RE.sub(" ", text or ""))).strip()


def round_cm(value: float) -> int:
    return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def unit_is_inch(unit: str | None) -> bool:
    return bool(unit and re.search(INCH_WORDS, unit, re.IGNORECASE))


def unit_is_cm(unit: str | None) -> bool:
    return bool(unit and re.search(CM_WORDS, unit, re.IGNORECASE))


def pair_from_text(text: str, source: str, assume_unitless_inches: bool) -> SizeInfo | None:
    for match in PAIR_RE.finditer(text or ""):
        a = float(match.group(1).replace(",", "."))
        b = float(match.group(2).replace(",", "."))
        unit = match.group(3)
        if unit_is_inch(unit):
            return SizeInfo("exact", round_cm(a * 2.54), round_cm(b * 2.54), source=source)
        if unit_is_cm(unit):
            return SizeInfo("exact", round_cm(a), round_cm(b), source=source)
        # Shopify's English wall-decal option values are normally inches. Very large values are treated as cm.
        assumed_inches = assume_unitless_inches and max(a, b) <= 130
        if assumed_inches:
            return SizeInfo("exact", round_cm(a * 2.54), round_cm(b * 2.54), source=source, assumed_inches=True)
        return SizeInfo("exact", round_cm(a), round_cm(b), source=source)
    return None


def range_from_explicit_units(text: str, source: str) -> SizeInfo | None:
    values_cm: list[int] = []
    for match in SINGLE_UNIT_RE.finditer(text or ""):
        value = float(match.group(1).replace(",", "."))
        unit = match.group(2)
        values_cm.append(round_cm(value * 2.54 if unit_is_inch(unit) else value))
    values_cm = [v for v in values_cm if 1 <= v <= 1000]
    if not values_cm:
        return None
    lo, hi = min(values_cm), max(values_cm)
    if lo == hi:
        return SizeInfo("single", max_cm=hi, source=source)
    return SizeInfo("range", min_cm=lo, max_cm=hi, source=source)


def option_text(variant: dict[str, Any]) -> tuple[str, str]:
    all_values: list[str] = []
    non_size: list[str] = []
    for opt in variant.get("selectedOptions") or []:
        name = str(opt.get("name") or "")
        value = str(opt.get("value") or "").strip()
        if not value or value.lower() == "default title":
            continue
        all_values.append(value)
        if not re.search(r"size|dimension|measurement|mõõt|suurus|izmērs|dydis|размер|koko", name, re.IGNORECASE):
            if pair_from_text(value, "option", True) is None:
                non_size.append(value)
    return " / ".join(all_values), " / ".join(non_size)


def parse_variant_size(variant: dict[str, Any], product: dict[str, Any]) -> tuple[SizeInfo | None, str]:
    options, extra = option_text(variant)
    title = str(variant.get("title") or "")
    for source, value, assume in (
        ("selected option", options, True),
        ("variant title", title, True),
    ):
        if value and value.lower() != "default title":
            size = pair_from_text(value, source, assume)
            if size:
                return size, extra

    product_title = str(product.get("title") or "")
    description = strip_html(str(product.get("descriptionHtml") or ""))
    size = pair_from_text(product_title, "product title", False)
    if size:
        return size, extra
    # Prefer an explicit pair from description, otherwise a min/max range of explicit measurements.
    size = pair_from_text(description, "description", False)
    if size:
        return size, extra
    return range_from_explicit_units(description, "description range"), extra


def translation_map(product: dict[str, Any], lang: str) -> tuple[dict[str, str], set[str]]:
    values: dict[str, str] = {}
    outdated: set[str] = set()
    for item in product.get(TRANSLATION_ALIAS[lang]) or []:
        key = str(item.get("key") or "")
        value = str(item.get("value") or "")
        if value:
            values[key] = value
        if item.get("outdated"):
            outdated.add(key)
    return values, outdated


def convert_inches_to_cm(text: str) -> str:
    if not text:
        return ""

    def repl_pair(match: re.Match[str]) -> str:
        a = float(match.group(1).replace(",", "."))
        b = float(match.group(2).replace(",", "."))
        unit = match.group(3)
        if not unit_is_inch(unit):
            return match.group(0)
        return f"{round_cm(a * 2.54)} × {round_cm(b * 2.54)} cm"

    result = PAIR_RE.sub(repl_pair, text)

    def repl_single(match: re.Match[str]) -> str:
        value = float(match.group(1).replace(",", "."))
        unit = match.group(2)
        if not unit_is_inch(unit):
            return match.group(0)
        return f"{round_cm(value * 2.54)} cm"

    return SINGLE_UNIT_RE.sub(repl_single, result)


def strip_parent_size(title: str) -> str:
    title = convert_inches_to_cm(title or "")
    # Remove explicit cm dimension pairs/ranges from parent titles when there are several variants.
    explicit_pair = re.compile(
        r"\s*(?:[-–—|,:()]\s*)*\d+(?:[\.,]\d+)?\s*[x×XхХ]\s*\d+(?:[\.,]\d+)?"
        r"(?:\s*[-–—]\s*\d+(?:[\.,]\d+)?\s*[x×XхХ]\s*\d+(?:[\.,]\d+)?)?"
        r"\s*(?:cm|см)\b\s*\*?",
        re.IGNORECASE,
    )
    title = explicit_pair.sub(" ", title)
    # Remove a trailing unitless dimension pair, but leave item counts such as 225x.
    trailing_pair = re.compile(
        r"\s*(?:[-–—|,:()]\s*)+\d+(?:[\.,]\d+)?\s*[x×XхХ]\s*\d+(?:[\.,]\d+)?\s*\*?\s*$",
        re.IGNORECASE,
    )
    title = trailing_pair.sub("", title)
    title = re.sub(r"\s*[-–—|,:]+\s*$", "", title)
    return SPACE_RE.sub(" ", title).strip()


def append_exact_size(title: str, size: SizeInfo, lang: str) -> str:
    title = convert_inches_to_cm(title or "")
    label = size.label(lang)
    # Remove existing trailing size so that only one normalized cm value remains.
    title = strip_parent_size(title)
    if label:
        return f"{title}, {label}".strip(" ,")
    return title


def derive_color_lt(text: str) -> str:
    t = text.lower()
    for pattern, value in (
        (r"\bblack\b|\bjuod|\bmust\b|\bч[её]рн", "Juoda"),
        (r"\bwhite\b|\bbalt|\bvalge\b|\bбел", "Balta"),
        (r"\bgold\b|\bauks|\bkuld|\bзолот", "Auksinė"),
        (r"\bsilver\b|\bsidabr|\bhõbe|\bсереб", "Sidabrinė"),
        (r"\bblue\b|\bmėlyn|\bzil|\bsinine\b|\bсин", "Mėlyna"),
        (r"\bgreen\b|\bžal|\brohe|\bзел", "Žalia"),
        (r"\bpink\b|\brož|\broosa|\bроз", "Rožinė"),
        (r"\bred\b|\braud|\bpunane\b|\bкрасн", "Raudona"),
        (r"\bgrey\b|\bgray\b|\bpilk|\bhall\b|\bсер", "Pilka"),
        (r"\bbrown\b|\brud|\bpruun|\bкорич", "Ruda"),
        (r"\byellow\b|\bgelton|\bkollane\b|\bжелт", "Geltona"),
    ):
        if re.search(pattern, t, re.IGNORECASE):
            return value
    return "Įvairiaspalvė"


def derive_pattern_lt(text: str) -> str:
    t = text.lower()
    if re.search(r"quote|saying|word|text|letter|inspir|motiv|family rules|bible|tekst|цитат|надпис", t):
        return "Tekstai"
    if re.search(r"star|moon|space|planet|zodiac|constellation|täht|kuu|звезд|луна", t):
        return "Žvaigždės"
    if re.search(r"flower|floral|leaf|tree|botanical|plant|gėl|zied|lill|цвет|лист", t):
        return "Gėlės"
    if re.search(r"map|world|globe|žemėlap|kaart|карта", t):
        return "Žemėlapiai"
    if re.search(r"line|stripe|dot|circle|geometric|triangle|hexagon|linij|joon|геометр|полос", t):
        return "Geometriniai"
    if re.search(r"animal|cat|dog|lion|tiger|bear|fox|elephant|giraffe|bird|butterfl|dinosaur|gyvū|loom|живот", t):
        return "Gyvūnai"
    return "Dekoratyviniai"


def market_prices(variant: dict[str, Any]) -> dict[str, tuple[Decimal, Decimal]] | None:
    result: dict[str, tuple[Decimal, Decimal]] = {}
    for lang in ("lt", "lv", "ee", "fi"):
        block = variant.get(lang) or {}
        current = block.get("price") or {}
        compare = block.get("compareAtPrice") or {}
        price = decimal(current.get("amount"))
        compare_price = decimal(compare.get("amount"))
        if price is None or str(current.get("currencyCode")) != "EUR":
            return None
        if compare_price is not None and str(compare.get("currencyCode")) == "EUR" and compare_price > price:
            regular, discounted = compare_price, price
        else:
            regular, discounted = price, price
        result[lang] = (q2(regular), q2(discounted))
    return result


def package_values(size: SizeInfo) -> tuple[Decimal, Decimal]:
    if size.kind == "exact" and size.width_cm and size.height_cm:
        short_side = min(size.width_cm, size.height_cm)
        long_side = max(size.width_cm, size.height_cm)
        area = Decimal(size.width_cm * size.height_cm) / Decimal(10000)
        tube_cm = min(max(short_side + 10, 25), 110)
        weight = max(Decimal("0.20"), min(Decimal("2.00"), Decimal("0.15") + area * Decimal("0.35")))
        return q2(weight), q2(Decimal(tube_cm) / Decimal(100))
    if size.kind == "range" and size.max_cm:
        length = min(max(size.max_cm + 10, 25), 110)
        return Decimal("0.35"), q2(Decimal(length) / Decimal(100))
    return Decimal("0.30"), Decimal("0.35")


def localized_size_list(variants: list[VariantRecord], lang: str) -> str:
    labels: list[str] = []
    seen: set[str] = set()
    for v in variants:
        label = v.size.label(lang)
        if v.extra_label:
            label = f"{v.extra_label} — {label}" if label else v.extra_label
        if label and label not in seen:
            seen.add(label)
            labels.append(label)
    return SIZE_INTRO[lang].format(items=", ".join(labels))


def normalize_description(base: str, variants: list[VariantRecord], lang: str) -> str:
    base = convert_inches_to_cm(base or "")
    # Remove scripts/styles and malformed null characters but preserve ordinary HTML.
    base = re.sub(r"<script\b[^>]*>.*?</script>", "", base, flags=re.IGNORECASE | re.DOTALL)
    base = re.sub(r"<style\b[^>]*>.*?</style>", "", base, flags=re.IGNORECASE | re.DOTALL)
    base = base.replace("\x00", "")
    sizes = localized_size_list(variants, lang)
    appendix = INSTALL_BLOCKS[lang].format(sizes=html.escape(sizes))
    return f"{base}{appendix}" if base else appendix


def make_product_record(product: dict[str, Any]) -> tuple[ProductRecord | None, str]:
    pid = numeric_product_id(str(product.get("id") or ""))
    if pid in IMPORTED_PRODUCT_IDS:
        return None, "already imported batch 001"
    title = str(product.get("title") or "").strip()
    handle = str(product.get("handle") or "").strip()
    if not title or not handle:
        return None, "missing title or handle"
    if CUSTOM_WORDS.search(title):
        return None, "custom/personalized product excluded"

    image_urls = [
        str(img.get("image", {}).get("url") or "")
        for img in product.get("_images", [])
        if img.get("image", {}).get("url")
    ]
    if not image_urls:
        return None, "no product image"

    raw_variants = sorted(product.get("_variants", []), key=lambda x: int(x.get("position") or 0))
    if not raw_variants:
        return None, "no variants"
    if len(raw_variants) >= 50:
        return None, "50 or more variants; export may be truncated"

    variants: list[VariantRecord] = []
    for raw_v in raw_variants:
        sku = str(raw_v.get("sku") or "").strip()
        ean = re.sub(r"\D", "", str(raw_v.get("barcode") or ""))
        if not sku:
            return None, "one or more variants missing SKU"
        if not valid_ean13(ean):
            return None, "one or more variants missing valid EAN-13"
        size, extra = parse_variant_size(raw_v, product)
        if size is None or not size.label("lt"):
            return None, "one or more variants have no reliably parsed size"
        prices = market_prices(raw_v)
        if prices is None:
            return None, "one or more variants missing EUR price for LT/LV/EE/FI"
        weight, length = package_values(size)
        variants.append(VariantRecord(raw_v, sku, ean, size, extra, prices, weight, length))

    localized_titles: dict[str, str] = {}
    localized_descriptions: dict[str, str] = {}
    missing: list[str] = []
    outdated: list[str] = []
    is_multi = len(variants) > 1
    for lang in LANGS:
        values, outdated_keys = translation_map(product, lang)
        tr_title = str(values.get("title") or "").strip()
        tr_body = str(values.get("body_html") or "").strip()
        if not tr_title:
            missing.append(f"{lang}:title")
            tr_title = title
        if not tr_body:
            missing.append(f"{lang}:description")
            tr_body = str(product.get("descriptionHtml") or "")
        if "title" in outdated_keys:
            outdated.append(f"{lang}:title")
        if "body_html" in outdated_keys:
            outdated.append(f"{lang}:description")

        if is_multi:
            localized_titles[lang] = strip_parent_size(tr_title)
        else:
            localized_titles[lang] = append_exact_size(tr_title, variants[0].size, lang)
        localized_descriptions[lang] = normalize_description(tr_body, variants, lang)

    confidence = "high"
    if any(v.size.assumed_inches for v in variants):
        confidence = "medium — unitless Shopify option interpreted as inches"
    if any(v.size.kind != "exact" for v in variants):
        confidence = "medium — size range/single measurement derived from description"

    combined = " ".join([title, strip_html(str(product.get("descriptionHtml") or ""))])
    return ProductRecord(
        raw=product,
        product_id=pid,
        title=title,
        handle=handle,
        images=image_urls[:5],
        variants=variants,
        localized_titles=localized_titles,
        localized_descriptions=localized_descriptions,
        missing_translations=missing,
        outdated_translations=outdated,
        size_confidence=confidence,
        color_lt=derive_color_lt(combined),
        pattern_lt=derive_pattern_lt(combined),
    ), ""


def cdata_el(parent: etree._Element, tag: str, value: Any) -> etree._Element:
    el = etree.SubElement(parent, tag)
    el.text = etree.CDATA(str(value))
    return el


def add_property(parent: etree._Element, prop_id: str, values: Iterable[Any]) -> None:
    p = etree.SubElement(parent, "property")
    cdata_el(p, "id", prop_id)
    vals = etree.SubElement(p, "values")
    for value in values:
        cdata_el(vals, "value", value)


def modification_title(v: VariantRecord, lang: str) -> str:
    size_label = v.size.label(lang)
    if v.extra_label:
        return f"{v.extra_label} — {size_label}" if size_label else v.extra_label
    return size_label


def product_to_xml(record: ProductRecord) -> etree._Element:
    p = etree.Element("product")
    cdata_el(p, "category-id", CATEGORY_ID)
    cdata_el(p, "category-name", CATEGORY_NAME)
    for lang in LANGS:
        cdata_el(p, TITLE_TAG[lang], record.localized_titles[lang])
    for lang in LANGS:
        cdata_el(p, DESC_TAG[lang], record.localized_descriptions[lang])

    cdata_el(p, "manufacturer-name", MANUFACTURER_NAME)
    cdata_el(p, "manufacturer-address", MANUFACTURER_ADDRESS)
    cdata_el(p, "manufacturer-email", MANUFACTURER_EMAIL)
    cdata_el(p, "composition", "Vinilas")
    cdata_el(
        p,
        "usage-info",
        "Išsami montavimo, priežiūros ir nuėmimo instrukcija pateikta prekės aprašyme. Klijuokite tik ant švaraus, lygaus ir visiškai sauso paviršiaus.",
    )

    props = etree.SubElement(p, "properties")
    add_property(props, "Prekės ženklas", ["Decords"])
    add_property(props, "Gamintojas", [MANUFACTURER_NAME])
    add_property(props, "Tipas", ["Sienų lipdukai"])
    add_property(props, "Medžiaga", ["Vinilas"])
    add_property(props, "Spalva", [record.color_lt])
    add_property(props, "Raštai", [record.pattern_lt])
    widths = sorted({v.size.width_cm for v in record.variants if v.size.kind == "exact" and v.size.width_cm})
    heights = sorted({v.size.height_cm for v in record.variants if v.size.kind == "exact" and v.size.height_cm})
    if widths:
        add_property(props, "Plotis", widths)
    if heights:
        add_property(props, "Aukštis", heights)

    colours = etree.SubElement(p, "colours")
    colour = etree.SubElement(colours, "colour")
    cdata_el(colour, "colour-title", "")
    images = etree.SubElement(colour, "images")
    for url in record.images:
        image = etree.SubElement(images, "image")
        cdata_el(image, "url", url)

    mods = etree.SubElement(colour, "modifications")
    for v in record.variants:
        m = etree.SubElement(mods, "modification")
        for lang in LANGS:
            cdata_el(m, MOD_TAG[lang], modification_title(v, lang))
        etree.SubElement(m, "weight").text = f"{v.weight_kg:.2f}"
        etree.SubElement(m, "length").text = f"{v.length_m:.2f}"
        etree.SubElement(m, "height").text = f"{v.height_m:.2f}"
        etree.SubElement(m, "width").text = f"{v.width_m:.2f}"
        attrs = etree.SubElement(m, "attributes")
        barcodes = etree.SubElement(attrs, "barcodes")
        cdata_el(barcodes, "barcode", v.ean)
        cdata_el(attrs, "supplier-code", v.sku)
        cdata_el(attrs, "manufacturer-code", v.sku)
    return p


def price_row_to_xml(parent: etree._Element, row: dict[str, Any]) -> None:
    p = etree.SubElement(parent, "product")
    order = (
        "sku", "ean",
        "price_lt", "price_after_discount_lt",
        "price_lv", "price_after_discount_lv",
        "price_ee", "price_after_discount_ee",
        "price_fi", "price_after_discount_fi",
        "stock",
        "collection_hours_lt", "collection_hours_lv", "collection_hours_ee", "collection_hours_fi",
    )
    for key in order:
        etree.SubElement(p, key).text = str(row[key])


def variant_price_row(v: VariantRecord) -> dict[str, Any]:
    row: dict[str, Any] = {"sku": v.sku, "ean": v.ean}
    for market in ("lt", "lv", "ee", "fi"):
        regular, discounted = v.prices[market]
        row[f"price_{market}"] = f"{regular:.2f}"
        row[f"price_after_discount_{market}"] = f"{discounted:.2f}"
    row["stock"] = STOCK
    for market in ("lt", "lv", "ee", "fi"):
        row[f"collection_hours_{market}"] = PRODUCTION_HOURS
    return row


def load_bulk_jsonl(url: str, destination: Path) -> dict[str, dict[str, Any]]:
    with requests.get(url, stream=True, timeout=180) as response:
        response.raise_for_status()
        with destination.open("wb") as f:
            for chunk in response.iter_content(chunk_size=1024 * 1024):
                if chunk:
                    f.write(chunk)

    products: dict[str, dict[str, Any]] = {}
    pending: dict[str, list[dict[str, Any]]] = defaultdict(list)
    with destination.open("r", encoding="utf-8") as f:
        for line_number, line in enumerate(f, 1):
            if not line.strip():
                continue
            obj = json.loads(line)
            parent_id = obj.pop("__parentId", None)
            if parent_id is None and "handle" in obj:
                obj["_images"] = []
                obj["_variants"] = []
                products[obj["id"]] = obj
                if obj["id"] in pending:
                    for child in pending.pop(obj["id"]):
                        if "sku" in child or "barcode" in child:
                            obj["_variants"].append(child)
                        elif "image" in child:
                            obj["_images"].append(child)
                continue
            if parent_id:
                product = products.get(parent_id)
                if product is None:
                    pending[parent_id].append(obj)
                elif "sku" in obj or "barcode" in obj:
                    product["_variants"].append(obj)
                elif "image" in obj:
                    product["_images"].append(obj)
    if pending:
        # Nested children without a known product should not occur, but do not silently ignore the condition.
        print(f"Warning: {len(pending)} parent IDs had orphan children", file=sys.stderr)
    return products


def parse_existing_price_rows(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    tree = etree.parse(str(path))
    rows: list[dict[str, Any]] = []
    for product in tree.findall(".//product"):
        row = {child.tag: (child.text or "") for child in product}
        rows.append(row)
    return rows


def write_xlsx(rows: list[dict[str, Any]], path: Path) -> None:
    headers = [
        "sku", "ean",
        "price_lt", "price_after_discount_lt",
        "price_lv", "price_after_discount_lv",
        "price_ee", "price_after_discount_ee",
        "price_fi", "price_after_discount_fi",
        "stock",
        "collection_hours_lt", "collection_hours_lv", "collection_hours_ee", "collection_hours_fi",
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "Feed"
    ws.append(headers)
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    ws.freeze_panes = "A2"
    for col_idx, header in enumerate(headers, 1):
        width = 18 if header not in {"stock"} else 12
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=2):
        for cell in row:
            cell.number_format = "@"
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=10):
        for cell in row:
            cell.number_format = "0.00"
    wb.save(path)


def write_report_csv(selected: list[ProductRecord], excluded: list[dict[str, str]], path: Path) -> None:
    headers = [
        "status", "shopify_product_id", "handle", "source_title", "variants", "first_sku", "first_ean",
        "size_confidence", "missing_translations", "outdated_translations", "reason",
    ]
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        for p in selected:
            writer.writerow({
                "status": "included",
                "shopify_product_id": p.product_id,
                "handle": p.handle,
                "source_title": p.title,
                "variants": len(p.variants),
                "first_sku": p.variants[0].sku,
                "first_ean": p.variants[0].ean,
                "size_confidence": p.size_confidence,
                "missing_translations": "; ".join(p.missing_translations),
                "outdated_translations": "; ".join(p.outdated_translations),
                "reason": "",
            })
        for row in excluded:
            writer.writerow({
                "status": "excluded",
                "shopify_product_id": row.get("product_id", ""),
                "handle": row.get("handle", ""),
                "source_title": row.get("title", ""),
                "variants": row.get("variants", ""),
                "first_sku": "",
                "first_ean": "",
                "size_confidence": "",
                "missing_translations": "",
                "outdated_translations": "",
                "reason": row.get("reason", ""),
            })


def write_report_xlsx(selected: list[ProductRecord], excluded: list[dict[str, str]], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Included products"
    headers = [
        "Batch", "Shopify product ID", "Handle", "Source title", "Variants", "First SKU", "First EAN",
        "LT parent title", "EE parent title", "RU parent title", "Size confidence",
        "Missing translations", "Outdated translations", "Package dimensions",
    ]
    ws.append(headers)
    for p in selected:
        ws.append([
            "002", p.product_id, p.handle, p.title, len(p.variants), p.variants[0].sku, p.variants[0].ean,
            p.localized_titles["lt"], p.localized_titles["ee"], p.localized_titles["ru"], p.size_confidence,
            "; ".join(p.missing_translations), "; ".join(p.outdated_translations), "Estimated",
        ])
    ex = wb.create_sheet("Excluded products")
    ex_headers = ["Shopify product ID", "Handle", "Source title", "Variants", "Reason"]
    ex.append(ex_headers)
    for row in excluded:
        ex.append([row.get("product_id", ""), row.get("handle", ""), row.get("title", ""), row.get("variants", ""), row.get("reason", "")])
    rules = wb.create_sheet("Rules")
    rule_rows = [
        ["Rule", "Value"],
        ["Multi-variant parent title", "No size and no size range"],
        ["Modification title", "Exact size in centimetres, rounded to whole centimetres"],
        ["Single fixed-size parent title", "Exact size in centimetres"],
        ["Stock", STOCK],
        ["Collection/production time", f"{PRODUCTION_HOURS} hours / 5 days"],
        ["Prices", "Current Shopify Markets EUR prices for LT, LV, EE and FI"],
        ["Discount", "Compare-at price is regular price only when it is higher than the current price"],
        ["Duplicate control", "Shopify product ID, SKU and EAN"],
        ["Content feed", "Cumulative: existing 11 products unchanged + new batch"],
        ["Existing content", "Do not click Content update"],
    ]
    for row in rule_rows:
        rules.append(row)

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        for col in range(1, sheet.max_column + 1):
            max_len = min(70, max(len(str(sheet.cell(r, col).value or "")) for r in range(1, min(sheet.max_row, 200) + 1)) + 2)
            sheet.column_dimensions[get_column_letter(col)].width = max(12, max_len)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--bulk-url-file", type=Path, required=True)
    parser.add_argument("--existing-products", type=Path, required=True)
    parser.add_argument("--existing-prices", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--target", type=int, default=TARGET_PRODUCTS)
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    bulk_url = args.bulk_url_file.read_text(encoding="utf-8").strip()
    if not bulk_url.startswith("https://"):
        raise ValueError("Bulk URL file does not contain an HTTPS URL")

    jsonl_path = args.output_dir / "shopify_wall_decals_source.jsonl"
    products = load_bulk_jsonl(bulk_url, jsonl_path)
    print(f"Loaded {len(products)} Shopify product roots")

    selected_strict: list[ProductRecord] = []
    selected_fallback: list[ProductRecord] = []
    excluded: list[dict[str, str]] = []
    seen_skus: set[str] = set()
    seen_eans: set[str] = set()

    existing_price_rows = parse_existing_price_rows(args.existing_prices)
    for row in existing_price_rows:
        if row.get("sku"):
            seen_skus.add(str(row["sku"]))
        if row.get("ean"):
            seen_eans.add(str(row["ean"]))

    for gid in sorted(products, key=lambda x: int(numeric_product_id(x))):
        raw = products[gid]
        record, reason = make_product_record(raw)
        if record is None:
            excluded.append({
                "product_id": numeric_product_id(gid),
                "handle": str(raw.get("handle") or ""),
                "title": str(raw.get("title") or ""),
                "variants": str(len(raw.get("_variants", []))),
                "reason": reason,
            })
            continue
        record_skus = {v.sku for v in record.variants}
        record_eans = {v.ean for v in record.variants}
        if record_skus & seen_skus:
            excluded.append({"product_id": record.product_id, "handle": record.handle, "title": record.title, "variants": str(len(record.variants)), "reason": "duplicate SKU"})
            continue
        if record_eans & seen_eans:
            excluded.append({"product_id": record.product_id, "handle": record.handle, "title": record.title, "variants": str(len(record.variants)), "reason": "duplicate EAN"})
            continue
        # Reserve identifiers only after the record passes duplicate checks.
        seen_skus.update(record_skus)
        seen_eans.update(record_eans)
        if record.missing_translations:
            selected_fallback.append(record)
        else:
            selected_strict.append(record)

    selected = (selected_strict + selected_fallback)[: args.target]
    selected_ids = {p.product_id for p in selected}
    for record in (selected_strict + selected_fallback)[args.target:]:
        excluded.append({"product_id": record.product_id, "handle": record.handle, "title": record.title, "variants": str(len(record.variants)), "reason": "eligible but outside current 500-product batch"})

    if not selected:
        raise RuntimeError("No eligible products were selected")

    batch_root = etree.Element("products", nsmap={"xsi": "http://www.w3.org/2001/XMLSchema-instance"})
    for record in selected:
        batch_root.append(product_to_xml(record))
    batch_path = args.output_dir / "PiguProducts-batch-002-new-only.xml"
    etree.ElementTree(batch_root).write(str(batch_path), encoding="UTF-8", xml_declaration=True, standalone=True, pretty_print=True)

    existing_tree = etree.parse(str(args.existing_products), etree.XMLParser(remove_blank_text=True, strip_cdata=False))
    master_root = etree.Element("products", nsmap={"xsi": "http://www.w3.org/2001/XMLSchema-instance"})
    for existing in existing_tree.findall(".//product"):
        master_root.append(deepcopy(existing))
    for record in selected:
        master_root.append(product_to_xml(record))
    master_path = args.output_dir / "PiguProducts-master.xml"
    etree.ElementTree(master_root).write(str(master_path), encoding="UTF-8", xml_declaration=True, standalone=True, pretty_print=True)

    new_price_rows: list[dict[str, Any]] = []
    for record in selected:
        for variant in record.variants:
            new_price_rows.append(variant_price_row(variant))
    all_price_rows = existing_price_rows + new_price_rows

    prices_root = etree.Element("products")
    for row in all_price_rows:
        price_row_to_xml(prices_root, row)
    prices_xml_path = args.output_dir / "PiguPrices-master.xml"
    etree.ElementTree(prices_root).write(str(prices_xml_path), encoding="UTF-8", xml_declaration=True, pretty_print=True)

    prices_xlsx_path = args.output_dir / "PiguPrices-master.xlsx"
    write_xlsx(all_price_rows, prices_xlsx_path)
    new_prices_xlsx_path = args.output_dir / "PiguPrices-batch-002-new-only.xlsx"
    write_xlsx(new_price_rows, new_prices_xlsx_path)

    report_csv = args.output_dir / "PiguBatch002_report.csv"
    report_xlsx = args.output_dir / "PiguBatch002_registry.xlsx"
    write_report_csv(selected, excluded, report_csv)
    write_report_xlsx(selected, excluded, report_xlsx)

    summary = {
        "source_product_roots": len(products),
        "target_new_products": args.target,
        "selected_new_products": len(selected),
        "selected_products_with_complete_translations": sum(not p.missing_translations for p in selected),
        "selected_products_with_translation_fallbacks": sum(bool(p.missing_translations) for p in selected),
        "new_offers": len(new_price_rows),
        "existing_products_preserved": len(existing_tree.findall(".//product")),
        "existing_offers_preserved": len(existing_price_rows),
        "master_products": len(master_root.findall("product")),
        "master_offers": len(all_price_rows),
        "stock": STOCK,
        "collection_hours": PRODUCTION_HOURS,
        "excluded_records": len(excluded),
        "parent_multi_variant_size_rule": "no size in parent title",
        "modification_size_rule": "exact rounded centimetres",
        "generated_files": [p.name for p in (batch_path, master_path, prices_xml_path, prices_xlsx_path, new_prices_xlsx_path, report_csv, report_xlsx)],
    }
    summary_path = args.output_dir / "PiguBatch002_summary.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    # Final structural and identifier checks.
    assert len(etree.parse(str(batch_path)).findall(".//product")) == len(selected)
    assert len(etree.parse(str(master_path)).findall(".//product")) == summary["master_products"]
    assert len(etree.parse(str(prices_xml_path)).findall(".//product")) == summary["master_offers"]
    new_skus = [row["sku"] for row in new_price_rows]
    new_eans = [row["ean"] for row in new_price_rows]
    assert len(new_skus) == len(set(new_skus))
    assert len(new_eans) == len(set(new_eans))
    assert all(valid_ean13(ean) for ean in new_eans)
    assert all(row["stock"] == STOCK for row in new_price_rows)
    assert all(row[f"collection_hours_{m}"] == PRODUCTION_HOURS for row in new_price_rows for m in ("lt", "lv", "ee", "fi"))

    print(json.dumps(summary, ensure_ascii=False, indent=2))
    # Source JSONL is not committed; it can be re-created from Shopify.
    jsonl_path.unlink(missing_ok=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
